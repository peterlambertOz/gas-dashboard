"""
Update the daily gas demand forecast workbook with any new forecast CSVs.

Usage: pip install openpyxl, then run this script.

Assumptions made (check these before relying on the output):
- The workbook and CSVs live directly in FORECAST_DIR and are .xlsx / .csv.
- "Most recent" existing workbook is the one whose filename date parses to
  the latest date (year is inferred as the current year, or the previous
  year if that would put the file in the future).
- New tabs are inserted in date order among the existing gas_forecast_*
  tabs; any other tabs in the workbook keep their current position.
- The formatting template is always sheet gas_forecast_20260619, row 3.
- The column letters in the brief (C:AD, AJ:AK, AV) refer to the CSV's
  original column layout, before the blank row is inserted at the top.
- "Convert text to values" only touches cells that actually parse as a
  number; text such as headers or labels is left untouched.
"""

import csv
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

FORECAST_DIR = Path("c:/Users/Peter/Python/Data/Forecasts")
CUTOFF_DATE = date(2026, 7, 1)  # only pick up csvs strictly after this date
TEMPLATE_SHEET = "gas_forecast_20260619"
TEMPLATE_ROW = 3
COLUMNS_TO_DELETE = ["C:AD", "AJ:AK", "AV"]


MONTH_NAMES = (
    "January|February|March|April|May|June|July|August|September|October|November|December"
)


def parse_workbook_date(path: Path) -> date:
    # Matches "as of 16 July", ignoring anything that follows before ".xlsx"
    # (e.g. a stray character or a "(1)" suffix from a duplicate download).
    match = re.search(rf"as of (\d{{1,2}}) ({MONTH_NAMES})", path.name)
    if not match:
        raise ValueError(f"Filename does not match expected pattern: {path.name}")
    day, month = match.group(1), match.group(2)
    today = date.today()
    parsed = datetime.strptime(f"{day} {month} {today.year}", "%d %B %Y").date()
    if parsed > today:
        parsed = datetime.strptime(f"{day} {month} {today.year - 1}", "%d %B %Y").date()
    return parsed


def find_latest_workbook(directory: Path) -> Path:
    candidates = list(directory.glob("Daily gas demand forecast data as of *.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No forecast workbook found in {directory}")
    return max(candidates, key=parse_workbook_date)


def find_new_forecast_csvs(directory: Path, existing_tabs: set[str], cutoff: date) -> list[Path]:
    pattern = re.compile(r"^gas_forecast_(\d{8})\.csv$")
    matches = []
    for csv_path in directory.glob("gas_forecast_*.csv"):
        match = pattern.match(csv_path.name)
        if not match:
            continue
        tab_name = f"gas_forecast_{match.group(1)}"
        csv_date = datetime.strptime(match.group(1), "%Y%m%d").date()
        if csv_date > cutoff and tab_name not in existing_tabs:
            matches.append(csv_path)
    matches.sort(key=lambda p: p.name)
    return matches


def col_range_to_indices(col_range: str) -> tuple[int, int]:
    parts = col_range.split(":")
    start = column_index_from_string(parts[0])
    end = column_index_from_string(parts[-1])
    return start, end


def delete_columns(sheet: Worksheet, ranges: list[str]) -> None:
    # Delete from the highest column index down, so earlier deletions
    # don't shift the positions of ranges still to be deleted.
    bounds = [col_range_to_indices(r) for r in ranges]
    bounds.sort(key=lambda b: b[0], reverse=True)
    for start, end in bounds:
        sheet.delete_cols(start, end - start + 1)


def copy_csv_into_sheet(csv_path: Path, sheet: Worksheet) -> None:
    with csv_path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            sheet.append(row)


def convert_text_to_values(sheet: Worksheet) -> None:
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            if isinstance(cell.value, str):
                try:
                    cell.value = float(cell.value.strip())
                except ValueError:
                    continue  # leave headers / labels as text


def find_last_data_row(sheet: Worksheet) -> int:
    for row_idx in range(sheet.max_row, 0, -1):
        if any(cell.value is not None for cell in sheet[row_idx]):
            return row_idx
    return sheet.max_row


def apply_template_row_format(
    template_sheet: Worksheet, template_row: int, target_sheet: Worksheet, first_row: int, last_row: int
) -> None:
    for target_row in range(first_row, last_row + 1):
        for col_idx in range(1, template_sheet.max_column + 1):
            template_cell = template_sheet.cell(row=template_row, column=col_idx)
            target_cell = target_sheet.cell(row=target_row, column=col_idx)
            target_cell.font = copy(template_cell.font)
            target_cell.fill = copy(template_cell.fill)
            target_cell.border = copy(template_cell.border)
            target_cell.alignment = copy(template_cell.alignment)
            target_cell.number_format = template_cell.number_format


def reorder_forecast_tabs(workbook: openpyxl.Workbook) -> None:
    pattern = re.compile(r"^gas_forecast_\d{8}$")
    positions = [i for i, name in enumerate(workbook.sheetnames) if pattern.match(name)]
    names_in_date_order = sorted(name for name in workbook.sheetnames if pattern.match(name))
    for position, name in zip(positions, names_in_date_order):
        workbook._sheets[position] = workbook[name]


def main() -> None:
    workbook_path = find_latest_workbook(FORECAST_DIR)
    print(f"Using workbook: {workbook_path.name}")

    workbook = openpyxl.load_workbook(workbook_path)
    existing_tabs = set(workbook.sheetnames)

    new_csvs = find_new_forecast_csvs(FORECAST_DIR, existing_tabs, CUTOFF_DATE)
    if not new_csvs:
        print("No new forecast CSVs found. Nothing to do.")
        return

    template_sheet = workbook[TEMPLATE_SHEET]

    for csv_path in new_csvs:
        tab_name = csv_path.stem
        print(f"Adding tab: {tab_name}")
        sheet = workbook.create_sheet(title=tab_name)

        copy_csv_into_sheet(csv_path, sheet)
        delete_columns(sheet, COLUMNS_TO_DELETE)
        sheet.insert_rows(1, 1)
        convert_text_to_values(sheet)

        last_data_row = find_last_data_row(sheet)
        apply_template_row_format(template_sheet, TEMPLATE_ROW, sheet, 3, last_data_row)

    reorder_forecast_tabs(workbook)

    pattern = re.compile(r"^gas_forecast_\d{8}$")
    forecast_tabs = [name for name in workbook.sheetnames if pattern.match(name)]
    newest_tab = max(forecast_tabs)
    for name in workbook.sheetnames:
        workbook[name].sheet_view.tabSelected = name == newest_tab
    workbook.active = workbook.sheetnames.index(newest_tab)

    today = date.today()
    output_name = f"Daily gas demand forecast data as of {today.day} {today.strftime('%B')}.xlsx"
    output_path = FORECAST_DIR / output_name
    workbook.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
